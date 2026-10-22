"""Document reading and rule-based extraction (text, contacts, skills) on realistic edge cases."""

import pytest

from app import ml_utils
from app.ml_utils import DocumentError, extract_skills, extract_text, normalize_skill, parse_contacts
from tests.conftest import SAMPLE_RESUMES


def skills(text, vocab):
    """Names only from ``extract_skills`` (the tests don't care about confidences)."""
    return [name for name, _ in extract_skills(text, vocab)]


@pytest.mark.parametrize(
    "text, vocab, expected",
    [
        ("Led R&D initiatives", ["r"], []),
        ("Let's go and ship it", ["go"], []),
        ("Objective-C and C-suite exposure", ["c"], []),
        ("We program in programming languages", ["r", "go", "c"], []),
        ("Data analysis in R and Python", ["r", "python"], ["python", "r"]),
        ("Microservices in Go (golang)", ["go"], ["go"]),
        ("Go beyond targets every quarter.", ["go"], []),
        ("Shipped on time. Go-to person for releases.", ["go"], []),
        ("Go, Python, SQL", ["go"], ["go"]),
        ("Languages: Go and Rust", ["go"], ["go"]),
        ("Built ML models with sklearn", ["scikit-learn"], ["scikit learn"]),
        ("Systems work in C/C++17", ["c", "c++"], ["c", "c++"]),
        ("Built APIs in C# and .NET; ASP.NET MVC", ["c#", ".net"], [".net", "c#"]),
        ("machine-learning pipelines", ["machine learning"], ["machine learning"]),
        ("k8s clusters and Postgres DBA", ["kubernetes", "postgresql"], ["kubernetes", "postgresql"]),
        ("ReactJS and Node.js developer", ["react", "node.js"], ["node.js", "react"]),
        ("Python3 scripting, CI/CD with GitHub Actions", ["python", "ci/cd"], ["ci/cd", "python"]),
        ("Java and JavaScript", ["java"], ["java"]),
        ("JavaScript only", ["java"], []),
    ],
)
def test_skill_extraction(text, vocab, expected):
    assert skills(text, vocab) == expected


def test_skill_confidence_reflects_match_type():
    found = dict(extract_skills("k8s and python", ["kubernetes", "python"]))
    assert found["python"] > found["kubernetes"]


def test_normalize_skill():
    assert normalize_skill("  Machine-Learning ") == "machine learning"
    assert normalize_skill("CI/CD") == "ci/cd"


@pytest.mark.parametrize(
    "phone",
    ["+1 (555) 234-5678", "+91 98765 43210", "+91-9876543210", "9876543210", "+44 20 7946 0958"],
)
def test_phone_formats(phone):
    assert parse_contacts(f"Jane Doe\n{phone}\njane@doe.com")["phone"] == phone


def test_phone_ignores_year_ranges():
    assert parse_contacts("Jane Doe\nEngineer 2018 - 2023 at Acme")["phone"] is None


@pytest.mark.parametrize(
    "text, name",
    [
        ("Curriculum Vitae\nJohn Doe\njohn@x.com", "John Doe"),
        ("JOHN DOE\njohn@x.com", "John Doe"),
        ("José García\njose@x.com", "José García"),
        ("Resume\nsoftware engineer", None),
    ],
)
def test_name_heuristic(text, name):
    assert parse_contacts(text)["name"] == name


def test_txt_decoding_falls_back_to_cp1252(tmp_path):
    path = tmp_path / "r.txt"
    path.write_bytes("Café ’quoted’".encode("cp1252"))
    assert extract_text(path) == "Café ’quoted’"


def test_docx_extraction_includes_text():
    text = extract_text(SAMPLE_RESUMES / "2_Priya_Patel_ML_DataScientist.docx")
    assert "Priya Patel" in text and len(text) > 300


@pytest.mark.parametrize("name, data", [("x.pdf", b"%PDF-1.4 truncated garbage"), ("x.docx", b"PK\x03\x04 not a zip"), ("x.doc", b"old")])
def test_unreadable_documents_raise(tmp_path, name, data):
    path = tmp_path / name
    path.write_bytes(data)
    with pytest.raises(DocumentError):
        extract_text(path)


@pytest.mark.parametrize("head, ext, ok", [(b"%PDF-1.7", ".pdf", True), (b"hello", ".pdf", False), (b"PK\x03\x04", ".docx", True),
                                           (b"<html>", ".docx", False), (b"plain", ".txt", True), (b"a\x00b", ".txt", False)])
def test_file_signature(head, ext, ok):
    if ok:
        ml_utils.check_file_signature(head, ext)
    else:
        with pytest.raises(DocumentError):
            ml_utils.check_file_signature(head, ext)


def test_embeddings_are_none_without_model():
    assert ml_utils.embed_text("anything") is None
    assert ml_utils.embed_batch(["a", "b"]) == [None, None]


def test_embeddings_with_model_are_normalized(fake_model):
    vector = ml_utils.embed_text("python developer")
    assert len(vector) == ml_utils.EMBEDDING_DIM
    assert abs(sum(v * v for v in vector) - 1.0) < 1e-6


def _minimal_pdf(text: str) -> bytes:
    """A one-page PDF whose content stream draws ``text`` (enough for pypdf to extract)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


def _encrypted_pdf(tmp_path, user_password: str):
    """Write a test PDF encrypted with ``user_password`` (empty = owner-password only)."""
    import io

    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(_minimal_pdf("Jane Doe Python engineer"))))
    writer.encrypt(user_password=user_password, owner_password="owner-secret", algorithm="AES-128")
    path = tmp_path / "cv.pdf"
    with open(path, "wb") as fh:
        writer.write(fh)
    return path


def test_pdf_text_extraction(tmp_path):
    path = tmp_path / "cv.pdf"
    path.write_bytes(_minimal_pdf("Jane Doe Python engineer"))
    assert "Jane Doe Python engineer" in extract_text(path)


def test_pdf_with_owner_password_only_is_readable(tmp_path):
    assert "Jane Doe" in extract_text(_encrypted_pdf(tmp_path, user_password=""))


def test_pdf_with_user_password_is_rejected(tmp_path):
    with pytest.raises(DocumentError, match="Password-protected"):
        extract_text(_encrypted_pdf(tmp_path, user_password="open-sesame"))


def test_docx_header_and_footer_text_is_extracted(tmp_path):
    import docx

    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "Jane Doe | jane@x.com | +1 415 555 0100"
    document.sections[0].footer.paragraphs[0].text = "References available on request"
    document.add_paragraph("Experience")
    document.add_paragraph("Python engineer")
    path = tmp_path / "cv.docx"
    document.save(path)
    text = extract_text(path)
    assert text.splitlines()[0] == "Jane Doe | jane@x.com | +1 415 555 0100"  # header first
    assert text.splitlines()[-1] == "References available on request"
    assert parse_contacts(text)["email"] == "jane@x.com"


def test_utf16_text_file(tmp_path):
    path = tmp_path / "cv.txt"
    path.write_bytes("Zoë Müller\nPython engineer".encode("utf-16"))  # with BOM, like Notepad "Unicode"
    ml_utils.check_file_signature(path.read_bytes()[:8], ".txt")  # NUL bytes are fine after a BOM
    assert extract_text(path).startswith("Zoë Müller")
    with pytest.raises(DocumentError):
        ml_utils.check_file_signature(b"MZ\x00\x00bin", ".txt")


@pytest.mark.parametrize(
    "raw, expected",
    [
        ("<b>Senior</b> engineer<script>alert(1)</script>", "Senior engineer"),
        ("Salary <50k and >20k per year", "Salary <50k and >20k per year"),
        ("Needs C++ (a<b) skills", "Needs C++ (a<b) skills"),
        ("  many   spaces \n\n second line ", "many spaces\n\nsecond line"),
    ],
)
def test_clean_text(raw, expected):
    from app.text_utils import clean_text

    assert clean_text(raw) == expected
