"""Document parsing, rule-based extraction, sentence embeddings and candidate scoring.

Everything here is a pure function of its inputs (plus the lazily loaded
embedding model), which keeps it easy to unit test.

Scoring model (single implementation, used by both the dashboard summary and the
persisted rankings)::

    skill_score = |required ∩ found| / |required|
    semantic    = (cos(job, resume) + 1) / 2            # 0.5 when an embedding is missing
    final       = min(1, 0.70·skill_score + 0.30·semantic + 0.05·[skill_score ≥ 0.70])
    status      = "accepted" if final ≥ 0.60 else "rejected"
"""

from __future__ import annotations

import logging
import re
import threading
import time
import unicodedata
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np

from .config import settings

logger = logging.getLogger("ml")

EMBEDDING_DIM = 384
SKILL_WEIGHT = 0.70
SEMANTIC_WEIGHT = 0.30
COVERAGE_BONUS = 0.05
BONUS_MIN_SKILL_SCORE = 0.70
ACCEPT_THRESHOLD = 0.60
NEUTRAL_SEMANTIC = 0.50
STATUS_ACCEPTED = "accepted"
STATUS_REJECTED = "rejected"


class DocumentError(ValueError):
    """The uploaded document cannot be read (wrong type, encrypted, corrupt, empty)."""


# =============================================================================
#  Embedding model (lazy, thread-safe, with retry back-off)
# =============================================================================
class _EmbeddingModel:
    """Holder for the sentence-transformer model, loaded lazily and thread-safely.

    Loading takes seconds (plus a ~90 MB download on first use), so it happens
    once per process, in a background thread at start-up. A failed load is
    retried at most every ``RETRY_AFTER_SECONDS`` instead of on every request.
    The name ``none`` (used by the tests) disables the model entirely.
    """

    RETRY_AFTER_SECONDS = 300

    def __init__(self, name: str) -> None:
        """Store the model name; nothing is loaded until ``get()`` is called."""
        self.name = name
        self._model = None
        self._lock = threading.Lock()
        self._failed_at: Optional[float] = None
        self.last_error: Optional[str] = None

    def get(self):
        """Return the loaded model, loading it on first use; ``None`` if unavailable.

        Double-checked locking: the fast path reads ``_model`` without the lock, and
        only the first caller pays for the import and download.
        """
        if self._model is not None:
            return self._model
        if not self.name or self.name.lower() == "none":
            return None
        if self._failed_at and time.monotonic() - self._failed_at < self.RETRY_AFTER_SECONDS:
            return None
        with self._lock:
            if self._model is None:
                try:
                    from sentence_transformers import SentenceTransformer

                    self._model = SentenceTransformer(self.name)
                    self._failed_at, self.last_error = None, None
                    logger.info("[ML] Loaded embedding model %s", self.name)
                except Exception as exc:  # missing package, no network, OOM…
                    self._failed_at = time.monotonic()
                    self.last_error = f"{exc.__class__.__name__}: {exc}"[:300]
                    logger.warning("[ML] Embedding model unavailable (%s); semantic scores will be neutral.", self.last_error)
        return self._model

    @property
    def loaded(self) -> bool:
        """True once the model is in memory (reported by ``/readyz`` and ``/metrics``)."""
        return self._model is not None


embedding_model = _EmbeddingModel(settings.embedding_model_name)


def embed_text(text: str) -> Optional[List[float]]:
    """Normalized 384-d embedding, or ``None`` if the model is unavailable.

    ``None`` is stored and later scored with a neutral semantic value instead of
    inventing a vector, so missing embeddings never masquerade as real ones.
    """
    return embed_batch([text])[0]


def embed_batch(texts: Sequence[str], batch_size: int = 32) -> List[Optional[List[float]]]:
    """Embed many texts in one model call (much faster than one call per text).

    Args:
        texts: Texts to embed; blank texts are replaced by a single space.
        batch_size: How many texts the model processes at a time.

    Returns:
        One unit-length 384-float list per text, in the same order, or ``None``
        for every text when the model is unavailable or fails.
    """
    if not texts:
        return []
    model = embedding_model.get()
    if model is None:
        return [None] * len(texts)
    clean = [t if t and t.strip() else " " for t in texts]
    try:
        vectors = model.encode(
            clean, batch_size=batch_size, show_progress_bar=False, convert_to_numpy=True, normalize_embeddings=True
        )
    except Exception:
        logger.exception("[ML] Embedding failed; storing NULL embeddings.")
        return [None] * len(texts)
    return [v.astype(float).tolist() for v in vectors]


# =============================================================================
#  File validation and text extraction
# =============================================================================
UTF16_BOMS = (b"\xff\xfe", b"\xfe\xff")  # Windows Notepad "Unicode" files start with one of these


def check_file_signature(head: bytes, extension: str) -> None:
    """Reject files whose content does not match their extension."""
    if extension == ".pdf" and not head.startswith(b"%PDF-"):
        raise DocumentError("The file is not a valid PDF document.")
    if extension == ".docx" and not head.startswith(b"PK\x03\x04"):
        raise DocumentError("The file is not a valid .docx document.")
    if extension == ".txt" and b"\x00" in head and not head.startswith(UTF16_BOMS):
        raise DocumentError("The .txt file contains binary data.")


def _docx_text(path: Path) -> str:
    """Extract text from a .docx file.

    Order: page headers first (contact details usually live there), then body
    paragraphs and table rows, then footers.

    Raises:
        DocumentError: The file is not a readable Word document.
    """
    try:
        import docx

        document = docx.Document(str(path))
    except Exception as exc:
        raise DocumentError("The .docx file could not be opened (corrupt or password-protected).") from exc

    def table_rows(tables):
        """Yield each table row as ``"cell | cell"`` with merged-cell repeats removed."""
        for table in tables:
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    value = cell.text.strip()
                    if value and value not in cells:  # merged cells repeat their text
                        cells.append(value)
                if cells:
                    yield " | ".join(cells)

    def block_text(block):
        """All non-empty paragraphs and table rows of a document, header or footer."""
        return [p.text for p in block.paragraphs if p.text.strip()] + list(table_rows(block.tables))

    # Many resumes keep the name, email and phone in the page header, so it goes first.
    headers, footers = [], []
    for section in document.sections:
        for header in (section.first_page_header, section.header):
            headers += block_text(header)
        for footer in (section.first_page_footer, section.footer):
            footers += block_text(footer)
    # Linked headers/footers repeat for every section; keep each line once.
    parts = list(dict.fromkeys(headers)) + block_text(document) + list(dict.fromkeys(footers))
    return "\n".join(parts)


def _pdf_text(path: Path) -> str:
    """Extract text from every page of a PDF.

    Encrypted PDFs are opened with an empty password first: many are locked only
    against editing or printing (owner password) and are otherwise readable.
    Scanned (image-only) PDFs yield little or no text; the caller rejects those.

    Raises:
        DocumentError: Password-protected, corrupt or unsupported PDF.
    """
    try:
        import pypdf

        reader = pypdf.PdfReader(str(path))
        if reader.is_encrypted:
            # Many PDFs are encrypted only with an owner password (printing/editing
            # restrictions) and open with an empty user password.
            try:
                opened = reader.decrypt("")
            except Exception:
                opened = 0
            if not opened:
                raise DocumentError("Password-protected PDFs are not supported.")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("The PDF could not be read (corrupt or unsupported).") from exc


def extract_text(path: Path) -> str:
    """Return the plain text of a .pdf, .docx or .txt file or raise ``DocumentError``."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".txt":
        data = path.read_bytes()
        if data.startswith(UTF16_BOMS):
            try:
                return data.decode("utf-16")
            except UnicodeDecodeError as exc:
                raise DocumentError("The .txt file is not valid UTF-16 text.") from exc
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode("latin-1")  # maps every byte, so it never fails
    if ext == ".docx":
        return _docx_text(path)
    if ext == ".pdf":
        return _pdf_text(path)
    raise DocumentError(f"Unsupported file type '{ext}'.")


# =============================================================================
#  Contact extraction
# =============================================================================
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}")
PHONE_CANDIDATE_RE = re.compile(r"(?<![\w])\+?\(?\d[\d\s().-]{7,18}\d(?![\w])")
NAME_STOPWORDS = {
    "resume", "curriculum", "vitae", "cv", "profile", "summary", "contact", "objective",
    "experience", "education", "skills", "projects", "page", "biodata",
    "engineer", "developer", "manager", "analyst", "scientist", "designer", "intern",
    "student", "consultant", "lead", "architect", "specialist", "fresher",
}


def _phone(text: str) -> Optional[str]:
    """First phone-like sequence with 10–15 digits, skipping year ranges like ``2019 - 2023``."""
    for match in PHONE_CANDIDATE_RE.finditer(text):
        candidate = match.group(0).strip()
        digits = re.sub(r"\D", "", candidate)
        if 10 <= len(digits) <= 15 and not re.fullmatch(r"(19|20)\d{2}\s*[-–]\s*(19|20)\d{2}", candidate):
            return candidate
    return None


def _name(lines: Iterable[str]) -> Optional[str]:
    """Guess the candidate's name from the first lines of a resume.

    A name line has 2–4 words, all capitalised ("Jane Doe") or upper-case
    ("JANE DOE"), no digits or ``@``, and no heading or job-title words
    (``NAME_STOPWORDS``). Upper-case names are returned in title case.
    """
    for line in list(lines)[:6]:
        if "@" in line or any(ch.isdigit() for ch in line) or len(line) > 60:
            continue
        cleaned = re.sub(r"[^\w\s.'-]", " ", line, flags=re.UNICODE)
        words = [w for w in cleaned.replace("_", " ").split() if w]
        if not 2 <= len(words) <= 4:
            continue
        if any(w.lower().strip(".") in NAME_STOPWORDS for w in words):
            continue
        # Names are capitalised ("Jane Doe") or written in capitals ("JANE DOE").
        if all(w[0].isalpha() and (w[0].isupper() or w.isupper()) for w in words):
            return " ".join(w.title() if w.isupper() else w for w in words)
    return None


def parse_contacts(text: str) -> Dict[str, Optional[str]]:
    """Best-effort ``{"name", "email", "phone"}`` from the top of a resume."""
    if not text:
        return {"name": None, "email": None, "phone": None}
    email = EMAIL_RE.search(text)
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    return {"name": _name(lines), "email": email.group(0).lower() if email else None, "phone": _phone(text)}


# =============================================================================
#  Skill extraction
# =============================================================================
# canonical skill -> extra spellings that count as the same skill
SKILL_ALIASES: Dict[str, Tuple[str, ...]] = {
    "kubernetes": ("k8s",),
    "postgresql": ("postgres",),
    "react": ("reactjs", "react.js"),
    "node.js": ("nodejs", "node js"),
    "ci/cd": ("ci cd", "cicd"),
    "rest api": ("restful api", "rest apis", "restful apis", "restful"),
    "aws": ("amazon web services",),
    "gcp": ("google cloud platform", "google cloud"),
    "nlp": ("natural language processing",),
    "c++": ("cpp",),
    "c#": ("csharp",),
    "scikit-learn": ("sklearn",),
    "javascript": ("ecmascript",),
}


def normalize_skill(name: str) -> str:
    """Canonical storage form: lowercase, single spaces, hyphens/underscores as spaces."""
    return re.sub(r"\s+", " ", re.sub(r"[-_]", " ", name.strip().lower()))


# Lookup table in canonical form ("scikit-learn" is stored as "scikit learn").
_ALIASES: Dict[str, Tuple[str, ...]] = {
    normalize_skill(skill): tuple(normalize_skill(a) for a in spellings) for skill, spellings in SKILL_ALIASES.items()
}

# One/two-letter skills that collide with ordinary words get case-sensitive rules.
_R_RE = re.compile(r"(?<![A-Za-z0-9&-])R(?![A-Za-z0-9&+#'-])")
_C_RE = re.compile(r"(?<![A-Za-z0-9&-])C(?![A-Za-z0-9&+#'-])")
_GO_RE = re.compile(r"(?<![A-Za-z0-9-])(Golang|golang|GO|Go)(?![A-Za-z0-9'-])")
_LIST_CONTEXT_RE = re.compile(r"[ \t]*(?:[,/;|)]|\n|$)")


def _mentions_go(text: str) -> bool:
    """True for "Golang", or for "Go"/"GO" used as a language rather than as the verb.

    "Go" opening a sentence ("Go beyond targets...") is the verb unless it is
    followed by list punctuation ("Go, Python, SQL").
    """
    for match in _GO_RE.finditer(text):
        if match.group(1).lower() == "golang":
            return True
        before = text[: match.start()].rstrip(" \t")
        opens_sentence = not before or before[-1] in ".!?\n"
        if not opens_sentence or _LIST_CONTEXT_RE.match(text, match.end()):
            return True
    return False


AMBIGUOUS_SKILL_MATCHERS = {
    "r": lambda text: _R_RE.search(text) is not None,
    "c": lambda text: _C_RE.search(text) is not None,
    "go": _mentions_go,
}


def _normalize_text(text: str) -> str:
    """Normalise resume text the same way skill names are stored.

    Lower-case, Unicode NFKC, hyphens/underscores as spaces and single spaces,
    so "Machine-Learning" in a resume matches the skill "machine learning".
    """
    text = unicodedata.normalize("NFKC", text).lower()
    return re.sub(r"\s+", " ", re.sub(r"[-_]", " ", text))


@lru_cache(maxsize=4096)
def _skill_regex(spelling: str) -> re.Pattern:
    """Compiled, cached pattern for one skill spelling with boundary rules.

    The guards stop partial-word matches ("java" inside "javascript", "sql"
    inside "mysql") while still allowing version suffixes ("python3", "c++17").
    """
    escaped = re.escape(spelling).replace(r"\ ", r"\s+")
    # Not preceded by a word char or symbol; not followed by a letter or symbol
    # (digits are allowed so "python3" and "C++17" still match).
    return re.compile(r"(?<![a-z0-9+#])" + escaped + r"(?![a-z+#])")


def extract_skills(text: str, vocabulary: Iterable[str]) -> List[Tuple[str, float]]:
    """Return ``[(canonical_skill, confidence)]`` for vocabulary skills present in ``text``."""
    if not text:
        return []
    normalized = _normalize_text(text)
    found: Dict[str, float] = {}
    for raw in vocabulary:
        skill = normalize_skill(raw or "")
        if not skill or skill in found:
            continue
        if skill in AMBIGUOUS_SKILL_MATCHERS:
            if AMBIGUOUS_SKILL_MATCHERS[skill](text):
                found[skill] = 0.9
            continue
        if _skill_regex(skill).search(normalized):
            found[skill] = 0.95
            continue
        if any(_skill_regex(alias).search(normalized) for alias in _ALIASES.get(skill, ())):
            found[skill] = 0.85
    return sorted(found.items())


# =============================================================================
#  Scoring
# =============================================================================
@dataclass(frozen=True)
class ResumeFeatures:
    """The two inputs scoring needs from a resume: its embedding and its skill names."""

    embedding: Optional[Sequence[float]]
    skills: Sequence[str]


def semantic_scores(job_embedding: Optional[Sequence[float]], resume_embeddings: Sequence[Optional[Sequence[float]]]) -> np.ndarray:
    """Vectorized ``(cos + 1) / 2`` for every resume; neutral 0.5 where a vector is missing."""
    n = len(resume_embeddings)
    scores = np.full(n, NEUTRAL_SEMANTIC, dtype=np.float64)
    if n == 0 or not job_embedding:
        return scores
    job = np.asarray(job_embedding, dtype=np.float32)
    job_norm = float(np.linalg.norm(job))
    valid = [i for i, v in enumerate(resume_embeddings) if v is not None and len(v) == job.shape[0]]
    if job_norm == 0.0 or not valid:
        return scores
    matrix = np.asarray([resume_embeddings[i] for i in valid], dtype=np.float32)  # (k, d)
    norms = np.linalg.norm(matrix, axis=1)
    norms[norms == 0] = 1.0  # avoid dividing by zero for an all-zero vector
    # One matrix-vector product gives every cosine at once: (k, d) @ (d,) -> (k,)
    cosines = (matrix @ job) / (norms * job_norm)
    scores[valid] = np.clip((cosines + 1.0) / 2.0, 0.0, 1.0)
    return scores


def score_candidates(
    job_embedding: Optional[Sequence[float]], job_skills: Sequence[str], resumes: Sequence[ResumeFeatures]
) -> List[Dict]:
    """Score many resumes against one job (one matrix-vector product for all cosines)."""
    required = [normalize_skill(s) for s in job_skills if s and s.strip()]
    required_unique = list(dict.fromkeys(required))
    required_set = set(required_unique)
    semantic = semantic_scores(job_embedding, [r.embedding for r in resumes])
    results = []
    for features, sem in zip(resumes, semantic, strict=True):
        have = {normalize_skill(s) for s in features.skills if s and s.strip()}
        if required_set:
            matched = required_set & have
            skill_score = len(matched) / len(required_set)
        else:
            # A job without required skills cannot reward or penalise skills: stay neutral.
            matched, skill_score = set(), NEUTRAL_SEMANTIC
        bonus = COVERAGE_BONUS if skill_score >= BONUS_MIN_SKILL_SCORE else 0.0
        final = min(1.0, SKILL_WEIGHT * skill_score + SEMANTIC_WEIGHT * float(sem) + bonus)
        coverage = [1 if s in have else 0 for s in required_unique]
        results.append(
            {
                "semantic": round(float(sem), 4),
                "skill_score": round(skill_score, 4),
                "final_score": round(final, 4),
                "status": STATUS_ACCEPTED if final >= ACCEPT_THRESHOLD else STATUS_REJECTED,
                "matched_skills": sorted(matched),
                "missing_skills": [s for s in required_unique if s not in have],
                "coverage_vector": coverage,
                "coverage_percent": int(round(100 * sum(coverage) / len(coverage))) if coverage else 0,
            }
        )
    return results

